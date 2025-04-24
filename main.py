import imaplib
import smtplib
import email
import json
from email.message import EmailMessage
from email.header import decode_header
from docx import Document
import os
import re
import tempfile
import secrets

# Настройки
IMAP_SERVER = "imap.hse.ru"
SMTP_SERVER = "smtp.hse.ru"
EMAIL_USER = secrets["email"] # "support@hse.ru"
EMAIL_PASS = secrets["password"] # "your_password_here"

# Подключение к IMAP для чтения писем
def get_emails():
    mail = imaplib.IMAP4_SSL(IMAP_SERVER)
    mail.login(EMAIL_USER, EMAIL_PASS)
    mail.select("inbox")

    # Получить только непрочитанные письма
    result, data = mail.search(None, 'UNSEEN')
    mail_ids = data[0].split()

    for num in mail_ids:
        result, data = mail.fetch(num, '(RFC822)')
        raw_email = data[0][1]
        msg = email.message_from_bytes(raw_email)

        # Декодирование темы письма
        subject = decode_header(msg["Subject"])[0][0]
        if isinstance(subject, bytes):
            subject = subject.decode()

        payload = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    charset = part.get_content_charset()
                    payload = part.get_payload(decode=True).decode(charset or 'utf-8')
                    break
        else:
            charset = msg.get_content_charset()
            payload = msg.get_payload(decode=True).decode(charset or 'utf-8')

        try:
            data = json.loads(payload)
            yield data
        except json.JSONDecodeError:
            print("Невозможно распарсить JSON")

    mail.logout()

# Создание документа по шаблону
def generate_docx_from_template(data, template_path="template.docx"):
    doc = Document(template_path)

    def replace_placeholders(text, values):
        def repl(match):
            key = match.group(1)
            return str(values.get(key, f"{{{{{key}}}}}"))
        return re.sub(r"{{(\w+)}}", repl, text)

    for para in doc.paragraphs:
        para.text = replace_placeholders(para.text, data)

    # Сохраняем во временный файл
    tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp_file.name)
    return tmp_file.name

# Отправка письма с вложением
def send_email(to_address, docx_path):
    msg = EmailMessage()
    msg["Subject"] = "Ваш проект договора подготовлен"
    msg["From"] = EMAIL_USER
    msg["To"] = to_address
    msg.set_content("Во вложении — документ, составленный на основе вашего письма.")

    with open(docx_path, "rb") as f:
        file_data = f.read()
        file_name = os.path.basename(docx_path)
        msg.add_attachment(file_data, maintype="application", subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename=file_name)

    with smtplib.SMTP_SSL(SMTP_SERVER) as server:
        server.login(EMAIL_USER, EMAIL_PASS)
        server.send_message(msg)

# Основной цикл
def main():
    for entry in get_emails():
        recipient = entry.get("email")  # предполагаем, что email указан в JSON
        if not recipient:
            print("Не указан адрес получателя.")
            continue
        doc_path = generate_docx(entry)
        send_email(recipient, doc_path)
        os.remove(doc_path)

if __name__ == "__main__":
    main()
